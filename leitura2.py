import pdfplumber
import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


VERBOSE = False

REGEX_CAMPOS = {
    'Nome': re.compile(
        r'10\s*-\s*Nome\s*\n\s*([A-ZÀ-Úa-zà-ú ]+)',
        re.IGNORECASE | re.DOTALL,
    ),
    'Número da Guia Principal': re.compile(
        r'3\s*-\s*Número da Guia Principal\s*\n?\s*\d+\s+(\d{5,})',
        re.IGNORECASE | re.DOTALL,
    ),
    'Número da Guia Principal Fallback': re.compile(
        r'2\s*-\s*N[°º]?\s*(?:da\s+)?Guia\s+no\s+Prestador\n.*?(\d{5,})',
        re.IGNORECASE | re.DOTALL,
    ),
    'Data Autorização': re.compile(
        r'4\s*-\s*Data da Autorização.*?\n\s*(\d{2}/\d{2}/\d{4})',
        re.IGNORECASE | re.DOTALL,
    ),
    'Senha': re.compile(
        r'4\s*-\s*Data da Autorização.*?\n\s*\d{2}/\d{2}/\d{4}\s+([A-Z0-9]+)',
        re.IGNORECASE | re.DOTALL,
    ),
    'Número da Carteira': re.compile(
        r'8\s*-\s*Número da Carteira.*?\n\s*(\d{5,})',
        re.IGNORECASE | re.DOTALL,
    ),
}

REGEX_PROCEDIMENTO = re.compile(
    r'(\d)\s+(\d{8})\s+(.+?)\s+(\d+)\s+(\d+)',
    re.IGNORECASE | re.DOTALL,
)
REGEX_CODIGO_TUSS = re.compile(r'\d{8}')

DESC_FORMATADA = {
    "TTO TEA E OUTROS TRANST GLOB DO DES-P DIA C/ FONO": "FONOAUDIOLOGIA",
    "PROCEDIMENTO PADRONIZADO PSICOPEDAGOGIA - POR DIA": "PSICOPEDAGOGIA",
    "TTO TEA E OUTROS TRANST GLOB DO DES-P DIA C/ PSICO": "PSICOTERAPIA",
    "TTO TEA E OUTROS TRANS GLOB DO DES-P DIA C/ TO": "TERAPIA OCUPACIONAL",
    "PP SEL ALIMEN TTO TEA E OUTROS TGD P/ DIA C/ NUTRI": "NUTRIÇÃO",
    "PP MUSICOTERAPIA - TEA E OUT TRANST GLOB DO DESENV": "MUSICOTERAPIA",
    "TTO TEA E OUTROS TRANS GLOB DO DES-P DIA C/ PSICOM": "PSICOMOTRICIDADE",
}


def extrair_dados_guia(page):
    """Extrai dados estruturados de uma página da guia TISS usando pdfplumber."""
    texto = page.extract_text() or ''
    dados = {}

    dados['Nome'] = extrair_campo(texto, REGEX_CAMPOS['Nome'])
    dados['Número da Guia Principal'] = extrair_campo(texto, REGEX_CAMPOS['Número da Guia Principal'])
    if not dados['Número da Guia Principal']:
        dados['Número da Guia Principal'] = extrair_campo(texto, REGEX_CAMPOS['Número da Guia Principal Fallback'])

    dados['Data Autorização'] = extrair_campo(texto, REGEX_CAMPOS['Data Autorização'])
    dados['Senha'] = extrair_campo(texto, REGEX_CAMPOS['Senha'])
    dados['Número da Carteira'] = extrair_campo(texto, REGEX_CAMPOS['Número da Carteira'])
    dados['Procedimentos'] = extrair_procedimentos_tabela(page, texto)

    return dados


def extrair_procedimentos_tabela(page, texto_fallback):
    """
    Tenta extrair procedimentos pelo texto já lido.
    Se não conseguir, faz fallback com tabela do pdfplumber.
    """
    procedimentos = []

    for match in REGEX_PROCEDIMENTO.finditer(texto_fallback):
        procedimentos.append({
            'Tabela': match.group(1),
            'Código do Procedimento': match.group(2),
            'Descrição': formatar_descricao(match.group(3)),
            'Qtde. Solic.': int(match.group(4)),
            'Qtde. Aut.': int(match.group(5)),
        })

    if procedimentos:
        return procedimentos

    if not REGEX_CODIGO_TUSS.search(texto_fallback):
        return procedimentos

    tabelas = page.extract_tables()
    for tabela in tabelas:
        for linha in tabela:
            if not linha:
                continue
            # Limpar células None
            celulas = [str(c).strip() if c else '' for c in linha]
            texto_linha = ' '.join(celulas)

            # Procurar linhas com código de 8 dígitos (padrão TUSS)
            match = REGEX_PROCEDIMENTO.search(texto_linha)
            if match:
                procedimentos.append({
                    'Tabela': match.group(1),
                    'Código do Procedimento': match.group(2),
                    'Descrição': formatar_descricao(match.group(3)),
                    'Qtde. Solic.': int(match.group(4)),
                    'Qtde. Aut.': int(match.group(5)),
                })

    return procedimentos


def extrair_campo(texto, pattern):
    """Helper para extrair um campo via regex."""
    match = pattern.search(texto)
    if match and match.lastindex:
        return match.group(1).strip()
    elif match:
        return match.group(0).strip()
    return ''


def exportar_para_excel(lista_dados, nome_arquivo='guias_extraidas.xlsx'):
    """Exporta dados das guias para Excel com formatação."""
    wb = Workbook()

    # Estilos
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws2 = wb.active
    ws2.title = 'Procedimentos'
    cols_proc = [
        'Nome',
        'N_Guia_Principal',
        'Dt_Autorizacao',
        'Senha',
        'N_Carteira',
        'Cod_Procedimento',
        'Descricao',
        'Qtde_Solic.',
        'Qtde_Aut.',
    ]
    ws2.append(cols_proc)
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for dados in lista_dados:
        for proc in dados.get('Procedimentos', []):
            ws2.append([
                dados.get('Nome'),
                dados.get('Número da Guia Principal', ''),
                dados.get('Data Autorização', ''),
                dados.get('Senha', ''),
                dados.get('Número da Carteira', ''),
                proc.get('Código do Procedimento', ''),
                proc.get('Descrição', ''),
                proc.get('Qtde. Solic.', ''),
                proc.get('Qtde. Aut.', ''),
            ])

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws2.column_dimensions[col[0].column_letter].width = min(max_len + 3, 50)

    wb.save(nome_arquivo)
    print(f'\nArquivo salvo: {nome_arquivo}')

def formatar_descricao(descricao_original: str) -> str:
    descricao_limpa = re.sub(r"\s+", " ", descricao_original).strip()
    return DESC_FORMATADA.get(descricao_limpa.upper(), descricao_limpa)

def processar_pdfs(pasta='assets'):
    arquivos = [f for f in os.listdir(pasta) if f.lower().endswith('.pdf')]
    todas_guias = []
    campos_contexto = [
        'Nome',
        'Número da Guia Principal',
        'Data Autorização',
        'Senha',
        'Número da Carteira',
    ]

    for nome_arquivo in arquivos:
        caminho = os.path.join(pasta, nome_arquivo)
        total_procedimentos = 0
        dados_ultima_pagina = {}
        print(f'\nProcessando: {nome_arquivo}')

        with pdfplumber.open(caminho) as pdf:
            for i, page in enumerate(pdf.pages):
                dados = extrair_dados_guia(page)
                for campo in campos_contexto:
                    if not dados.get(campo):
                        dados[campo] = dados_ultima_pagina.get(campo, '')

                if any(dados.get(campo) for campo in campos_contexto):
                    dados_ultima_pagina = {
                        campo: dados.get(campo, '')
                        for campo in campos_contexto
                    }

                dados['Arquivo Origem'] = nome_arquivo
                dados['Página'] = i + 1
                todas_guias.append(dados)
                total_procedimentos += len(dados.get('Procedimentos', []))

                if VERBOSE:
                    print(f'  Página {i+1}:')
                    for k, v in dados.items():
                        if k not in ('Procedimentos', 'Arquivo Origem', 'Página') and v:
                            print(f'    {k}: {v}')
                    for proc in dados.get('Procedimentos', []):
                        print(f'    -> {proc["Código do Procedimento"]} | {proc["Descrição"]} '
                              f'| Solic: {proc["Qtde. Solic."]} | Aut: {proc["Qtde. Aut."]}')

        print(f'  {len(pdf.pages)} pagina(s), {total_procedimentos} procedimento(s)')

    return todas_guias


if __name__ == '__main__':
    guias = processar_pdfs()
    if guias:
        exportar_para_excel(guias)
    else:
        print('Nenhuma guia encontrada nos PDFs.')
